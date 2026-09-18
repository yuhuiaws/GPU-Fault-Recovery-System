from __future__ import annotations

import hashlib
import sys
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook

from tools import generate_nvidia_xid_policy as generator


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    book = Workbook()
    xids = book.active
    xids.title = "Xids"
    xids.append([f"column-{index}" for index in range(12)])
    xids.append(
        [
            None,
            14,
            " TEST ",
            " Description ",
            "YES",
            "NO",
            "YES",
            None,
            " RESTART_APP ",
            None,
            "",
            " condition ",
        ]
    )
    decode = book.create_sheet("Xid 144-150 Decode")
    decode.append([f"column-{index}" for index in range(13)])
    decode.append(
        [
            144,
            " SUBCODE ",
            "0000",
            "0001",
            None,
            "IGNORE",
            "",
            None,
            None,
            "Inspect",
            "Warning",
            "GPU",
            "local",
        ]
    )
    buckets = book.create_sheet("Resolution Buckets")
    buckets.append(["guidance", "resolution"])
    buckets.append([" Restart ", " Recover "])
    buckets.append([None, "ignored"])
    path = tmp_path / "catalog.xlsx"
    book.save(path)
    book.close()
    return path


@pytest.mark.parametrize(
    "value,expected", [(None, None), (" ", None), (" name ", "name"), (42, "42")]
)
def test_catalog_text_normalizes_whitespace_without_inventing_values(
    value, expected
) -> None:
    assert generator.text(value) == expected, "empty cells remain absent"


@pytest.mark.parametrize("value", [7, 7.0, " 7 "])
def test_integral_catalog_ids_have_one_representation(value) -> None:
    assert generator.integer(value) == 7, (
        "numeric and textual spreadsheet cells normalize to the same ID"
    )


@pytest.mark.parametrize("value", [None, 7.5, "invalid", True, False])
def test_invalid_catalog_ids_cannot_be_coerced(value) -> None:
    with pytest.raises((TypeError, ValueError)):
        generator.integer(value)


def test_workbook_maps_policy_fields_and_closes_the_read_handle(
    workbook, monkeypatch
) -> None:
    closed = []
    load = generator.load_workbook

    def tracked(*args, **kwargs):
        book = load(*args, **kwargs)
        close = book.close

        def close_book():
            closed.append(True)
            close()

        book.close = close_book
        return book

    monkeypatch.setattr(generator, "load_workbook", tracked)
    document = generator.generate(workbook, "unit")
    assert (
        document["metadata"]["sourceSha256"]
        == hashlib.sha256(workbook.read_bytes()).hexdigest()
    ), "the policy retains exact source identity"
    assert document["spec"]["catalogRules"] == [
        {
            "xid": 14,
            "mnemonic": "TEST",
            "description": "Description",
            "products": ["A100", "B100"],
            "immediateAction": "RESTART_APP",
            "investigatoryAction": None,
            "xid154Linkage": None,
            "triggerConditions": "condition",
        }
    ], "column mappings preserve product applicability and action semantics"
    assert document["spec"]["nvlink5"]["decodeRules"][0]["subcodeName"] == "SUBCODE", (
        "decode fields use the same normalization"
    )
    assert document["spec"]["resolutionBuckets"] == {"Restart": "Recover"}, (
        "blank guidance rows do not create a bucket"
    )
    assert closed == [True], "a reusable generator must release the read-only workbook"


def test_missing_sheet_still_closes_the_workbook(tmp_path, monkeypatch) -> None:
    source = tmp_path / "incomplete.xlsx"
    source.write_bytes(b"unit")
    closed = []

    class MissingSheets:
        def __getitem__(self, name):
            raise KeyError(name)

        def close(self):
            closed.append(True)

    monkeypatch.setattr(
        generator, "load_workbook", lambda *args, **kwargs: MissingSheets()
    )
    with pytest.raises(KeyError):
        generator.generate(source, "unit")
    assert closed == [True], "parse failure cannot leak the workbook archive"


def test_render_round_trips_as_ascii_and_has_a_stable_digest(workbook) -> None:
    policy = generator.generate(workbook, "unit")
    policy["spec"]["catalogRules"][0]["description"] = "\u6d4b\u8bd5"
    first = generator.render(policy)
    assert first.startswith(generator.HEADER), "generated assets identify their source"
    first.encode("ascii")
    parsed = yaml.safe_load(first)
    assert parsed["spec"]["catalogRules"][0]["description"] == "\u6d4b\u8bd5", (
        "ASCII output preserves Unicode data through escaping"
    )
    assert generator.render(parsed) == first, "repeated rendering is canonical"
    assert len(parsed["metadata"][generator.GENERATED_SHA_FIELD]) == 64, (
        "the output includes its content identity"
    )


def test_committed_generated_policy_passes_full_integrity_checks() -> None:
    generator.check_generated(generator.DEFAULT_OUTPUT)


@pytest.mark.parametrize("problem", ["unicode", "noncanonical"])
def test_invalid_generated_asset_is_rejected(tmp_path, problem) -> None:
    path = tmp_path / "generated.yaml"
    original = generator.DEFAULT_OUTPUT.read_text(encoding="utf-8")
    path.write_text(
        original + ("\n# \u6d4b\u8bd5\n" if problem == "unicode" else "\n"),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="ASCII|canonical"):
        generator.check_generated(path)


def test_cli_generation_checks_source_digest_before_writing(
    workbook, tmp_path, monkeypatch
) -> None:
    output = tmp_path / "generated/policy.yaml"
    digest = hashlib.sha256(workbook.read_bytes()).hexdigest()
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "generator",
            str(workbook),
            str(output),
            "--catalog-version",
            "unit",
            "--expected-sha256",
            digest,
        ],
    )
    generator.main()
    assert yaml.safe_load(output.read_text())["metadata"]["sourceSha256"] == digest, (
        "explicit source identity is preserved"
    )
    before = output.read_bytes()
    monkeypatch.setattr(
        sys,
        "argv",
        ["generator", str(workbook), str(output), "--expected-sha256", "0" * 64],
    )
    with pytest.raises(SystemExit, match="digest mismatch"):
        generator.main()
    assert output.read_bytes() == before, (
        "failed source verification cannot overwrite the last artifact"
    )


@pytest.mark.parametrize("arguments", [[], ["--check", "unexpected.xlsx"]])
def test_cli_rejects_incomplete_or_conflicting_modes(monkeypatch, arguments) -> None:
    monkeypatch.setattr(sys, "argv", ["generator", *arguments])
    with pytest.raises(SystemExit) as error:
        generator.main()
    assert error.value.code == 2, "mode errors are rejected before workbook reads"


def test_check_cli_only_reads_the_selected_asset(monkeypatch, capsys, tmp_path) -> None:
    monkeypatch.setattr(sys, "argv", ["generator", "--check"])
    generator.main()
    assert "catalog is canonical" in capsys.readouterr().out, (
        "integrity check reports its validated inventory"
    )
    monkeypatch.setattr(
        sys,
        "argv",
        ["generator", "--check", "--generated-policy", str(tmp_path / "missing")],
    )
    with pytest.raises(SystemExit, match="catalog check failed"):
        generator.main()
