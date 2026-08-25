from __future__ import annotations

import argparse
import hashlib
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

from gpu_fault.policy.catalog_integrity import (
    EXPECTED_NVLINK5_DECODE_RULES,
    EXPECTED_RESOLUTION_BUCKETS,
    EXPECTED_XID_RULES,
    GENERATED_SHA_FIELD,
    PINNED_XID_SOURCE_SHA256,
    validate_xid_catalog_document,
    xid_catalog_artifact_sha256,
)


SOURCE_URL = (
    "https://docs.nvidia.com/deploy/xid-errors/"
    "_downloads/4586dadb59119a55d1e93a181caa4272/"
    "Xid-Catalog.xlsx"
)
ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "src/gpu_fault/data/nvidia-xid-catalog-610.generated.yaml"
HEADER = "# Generated from the NVIDIA Xid Catalog. Do not edit.\n"


def text(value: Any) -> str | None:
    if value is None:
        return None
    normalized = str(value).strip()
    return normalized or None


def integer(value: Any) -> int:
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return int(str(value).strip())


def generate(source: Path, catalog_version: str) -> dict[str, Any]:
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)

    xid_sheet = workbook["Xids"]
    catalog_rules = []
    for row in xid_sheet.iter_rows(min_row=2, values_only=True):
        products = [
            product
            for product, value in zip(
                ("A100", "H100", "B100", "GB200"),
                row[4:8],
                strict=True,
            )
            if text(value) == "YES"
        ]
        catalog_rules.append(
            {
                "xid": integer(row[1]),
                "mnemonic": text(row[2]),
                "description": text(row[3]),
                "products": products,
                "immediateAction": text(row[8]),
                "investigatoryAction": text(row[9]),
                "xid154Linkage": text(row[10]),
                "triggerConditions": text(row[11]),
            }
        )

    decode_sheet = workbook["Xid 144-150 Decode"]
    decode_rules = []
    for row in decode_sheet.iter_rows(min_row=2, values_only=True):
        decode_rules.append(
            {
                "xid": integer(row[0]),
                "subcodeName": text(row[1]),
                "v1Pattern": text(row[2]),
                "v2Pattern": text(row[3]),
                "errorStatus": text(row[4]),
                "recoveryAction": text(row[5]),
                "action2V1Pattern": text(row[6]),
                "action2V2Pattern": text(row[7]),
                "action2": text(row[8]),
                "investigatoryAction": text(row[9]),
                "severity": text(row[10]),
                "faultOrigin": text(row[11]),
                "localRemote": text(row[12]),
            }
        )

    bucket_sheet = workbook["Resolution Buckets"]
    resolution_buckets = {}
    for guidance, resolution in bucket_sheet.iter_rows(min_row=2, values_only=True):
        key = text(guidance)
        if key:
            resolution_buckets[key] = text(resolution)

    return {
        "apiVersion": "gpu-fault.io/v1alpha1",
        "kind": "NvidiaXidCatalogPolicy",
        "metadata": {
            "name": f"nvidia-xid-catalog-{catalog_version}",
            "generatedFrom": SOURCE_URL,
            "sourceSha256": digest,
        },
        "spec": {
            "catalog": {
                "version": catalog_version,
                "coverage": "FULL_OFFICIAL_ARTIFACT",
                "productFamilies": [
                    {"family": "A100", "modelPrefixes": ["A"]},
                    {"family": "H100", "modelPrefixes": ["H", "GH"]},
                    {"family": "B100", "modelPrefixes": ["B"]},
                    {"family": "GB200", "modelPrefixes": ["GB"]},
                ],
            },
            "defaults": {"markerTtlSeconds": 3600},
            "correlation": {"companionWindowSeconds": 30},
            "catalogRules": catalog_rules,
            "nvlink5": {
                "driverBoundary": 575,
                "decodeRules": decode_rules,
            },
            "resolutionBuckets": resolution_buckets,
        },
    }


def render(policy: dict[str, Any]) -> str:
    policy["metadata"][GENERATED_SHA_FIELD] = xid_catalog_artifact_sha256(policy)
    return HEADER + yaml.safe_dump(
        policy,
        sort_keys=False,
        allow_unicode=False,
        width=1000,
    )


def check_generated(path: Path) -> None:
    content = path.read_text(encoding="utf-8")
    try:
        content.encode("ascii")
    except UnicodeEncodeError as exc:
        raise ValueError("generated XID catalog must be ASCII") from exc
    document = yaml.safe_load(content)
    validate_xid_catalog_document(document)
    if render(document) != content:
        raise ValueError("generated XID catalog is not in canonical generator format")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path, nargs="?")
    parser.add_argument("output", type=Path, nargs="?")
    parser.add_argument("--catalog-version", default="610")
    parser.add_argument(
        "--expected-sha256",
        default=PINNED_XID_SOURCE_SHA256,
    )
    parser.add_argument("--check", action="store_true")
    parser.add_argument(
        "--generated-policy",
        type=Path,
        default=DEFAULT_OUTPUT,
    )
    args = parser.parse_args()

    if args.check:
        if args.source is not None or args.output is not None:
            parser.error("--check does not accept source/output arguments")
        try:
            check_generated(args.generated_policy)
        except (OSError, ValueError, yaml.YAMLError) as exc:
            raise SystemExit(f"generated XID catalog check failed: {exc}") from exc
        print(
            "Generated XID catalog is canonical: "
            f"{EXPECTED_XID_RULES} rules, "
            f"{EXPECTED_NVLINK5_DECODE_RULES} NVLink5 decodes, "
            f"{EXPECTED_RESOLUTION_BUCKETS} resolution buckets."
        )
        return
    if args.source is None:
        parser.error("source XLSX is required unless --check is used")
    output = args.output or DEFAULT_OUTPUT
    policy = generate(args.source, args.catalog_version)
    actual_digest = policy["metadata"]["sourceSha256"]
    if actual_digest != args.expected_sha256:
        raise SystemExit(f"catalog digest mismatch: {actual_digest}")

    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(render(policy), encoding="utf-8")


if __name__ == "__main__":
    main()
